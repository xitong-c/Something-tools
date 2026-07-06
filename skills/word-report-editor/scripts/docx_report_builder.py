from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


@dataclass
class FigureSpec:
    caption: str
    hint: str = "Insert figure here"


@dataclass
class SectionSpec:
    heading: str
    paragraphs: list[str] = field(default_factory=list)
    figures: list[FigureSpec] = field(default_factory=list)
    page_break_after: bool = False


def set_run_font(run, size: int = 12, bold: bool = False, east_asia: str = "宋体") -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold


def set_paragraph_format(paragraph, first_line: bool = True, line_spacing: float = 1.5) -> None:
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.space_after = Pt(6)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Pt(24)


def clear_document_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def add_paragraph(
    doc: Document,
    text: str = "",
    *,
    size: int = 12,
    bold: bool = False,
    align=None,
    first_line: bool = True,
) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    set_paragraph_format(p, first_line=first_line)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    style_name = "Heading 2" if level == 1 else "Heading 3"
    try:
        p = doc.add_paragraph(style=style_name)
    except KeyError:
        p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(p, first_line=False)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 14, bold=True, east_asia="黑体")


def add_toc_field(doc: Document, title: str = "目  录") -> None:
    add_paragraph(doc, title, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Open in Word/WPS and update this table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = p.add_run()
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(sep)
    run._r.append(text)
    run._r.append(end)
    set_run_font(run)


def add_placeholder_figure(doc: Document, spec: FigureSpec) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, first_line=False, line_spacing=1.2)
    run = p.add_run(f"\n\n{spec.hint}\n\n")
    set_run_font(run, size=11)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "8")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "999999")
        borders.append(tag)
    tc_pr.append(borders)
    add_paragraph(doc, spec.caption, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], caption: str) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        run = table.rows[0].cells[i].paragraphs[0].add_run(header)
        set_run_font(run, size=10, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            run = cells[i].paragraphs[0].add_run(value)
            set_run_font(run, size=10)
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                tag = OxmlElement(f"w:{edge}")
                tag.set(qn("w:val"), "single")
                tag.set(qn("w:sz"), "4")
                tag.set(qn("w:space"), "0")
                tag.set(qn("w:color"), "666666")
                borders.append(tag)
            tc_pr.append(borders)
    add_paragraph(doc, caption, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)


def build_report(
    template_path: str | Path,
    output_path: str | Path,
    title: str,
    report_name: str,
    sections: list[SectionSpec],
    cover_lines: list[str] | None = None,
) -> Path:
    doc = Document(str(template_path))
    clear_document_body(doc)
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.3)
        section.left_margin = Cm(2.6)
        section.right_margin = Cm(2.4)

    add_paragraph(doc, report_name, size=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_paragraph(doc, title, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    for line in cover_lines or []:
        add_paragraph(doc, line, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    doc.add_page_break()

    add_toc_field(doc)
    doc.add_page_break()

    add_heading(doc, title, level=1)
    for section in sections:
        add_heading(doc, section.heading, level=1)
        for paragraph in section.paragraphs:
            add_paragraph(doc, paragraph)
        for figure in section.figures:
            add_placeholder_figure(doc, figure)
        if section.page_break_after:
            doc.add_page_break()

    out = Path(output_path)
    doc.save(out)
    return out


if __name__ == "__main__":
    print("Import and call build_report() from a task-specific generation script.")
